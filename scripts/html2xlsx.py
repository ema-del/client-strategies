# Convert a Google Sheets "zip (HTML)" export into the same JSON+base64 xlsx file scale_data.py reads.
import sys, json, base64, zipfile, io, re
from html.parser import HTMLParser
import openpyxl

class T(HTMLParser):
    def __init__(s): super().__init__(); s.rows=[]; s.row=None; s.cell=None
    def handle_starttag(s,t,a):
        if t=='tr': s.row=[]
        elif t in('td','th') and s.row is not None:
            s.cell=''; s.span=int(dict(a).get('colspan','1') or 1)
        elif t=='br' and s.cell is not None: s.cell+='\n'
    def handle_endtag(s,t):
        if t in('td','th') and s.cell is not None:
            s.row.append(s.cell.strip()); s.row += ['']*(s.span-1); s.cell=None
        elif t=='tr' and s.row is not None: s.rows.append(s.row); s.row=None
    def handle_data(s,d):
        if s.cell is not None: s.cell+=d

def num(v):
    x=v.replace(',','').replace('£','').replace('$','').replace('€','').strip()
    if re.fullmatch(r'-?\d+(\.\d+)?%',x): return float(x[:-1])/100
    if re.fullmatch(r'-?\d+(\.\d+)?',x): return float(x)
    m2=re.fullmatch(r'[A-Z][a-z]{2},? ([A-Z][a-z]{2}) (\d{1,2})',x)
    if m2:
        import datetime; return datetime.datetime.strptime(f"{m2[1]} {m2[2]} {YEAR}",'%b %d %Y')
    m=re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})',x)
    if m:
        import datetime; return datetime.datetime(int(m[3]),int(m[1]),int(m[2]))
    return v or None

src, out = sys.argv[1], sys.argv[2]
YEAR = int(sys.argv[3]) if len(sys.argv)>3 else 2026
d=json.load(open(src)); z=zipfile.ZipFile(io.BytesIO(base64.b64decode(d['content'])))
wb=openpyxl.Workbook(); wb.remove(wb.active)
for n in z.namelist():
    if not n.endswith('.html') or '/' in n: continue
    p=T(); p.feed(z.read(n).decode('utf-8','replace'))
    ws=wb.create_sheet(n[:-5][:31])
    rows=[r[1:] for r in p.rows]                         # drop the row-number column
    rows=[r for r in rows if not all(re.fullmatch(r'[A-Z]{1,2}|',c or '') for c in r)]  # drop the A,B,C letter row
    width=max((len(r) for r in rows),default=0)
    keep=[i for i in range(width) if any(i<len(r) and r[i] for r in rows)]  # drop empty spacer columns
    for r in rows:
        ws.append([num(r[i]) if i<len(r) else None for i in keep])
buf=io.BytesIO(); wb.save(buf)
json.dump({'content':base64.b64encode(buf.getvalue()).decode(),'id':d.get('id'),'title':d.get('title'),'mimeType':'xlsx-from-html'},open(out,'w'))
print('ok', wb.sheetnames)
